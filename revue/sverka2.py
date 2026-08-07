# -*- coding: utf-8 -*-
"""РЕВЬЮ (без правок): сверка «Спецификация» (итоговый) ↔ сводный ВОР, с корректным разбором
текстовых чисел в ВОРе."""
import openpyxl, re, json
from collections import defaultdict
ITOG='/root/.claude/uploads/a6572ada-bc7d-590d-9f64-07c4882f0066/330c2672-SPEC_IOS_TH_ITOGOVIY_FILE2.xlsx'
VOR ='/root/.claude/uploads/a6572ada-bc7d-590d-9f64-07c4882f0066/f840d37a-___________.xlsx'
def s(x): return '' if x is None else re.sub(r'\s+',' ',str(x)).strip()
def key(t): return re.sub(r'[^а-яёa-z0-9]','', s(t).lower().replace('ё','е'))
def N(x):
    """число из ячейки: принимает и float, и текст «1 881», и «1,5»"""
    if isinstance(x,(int,float)): return float(x)
    t=s(x).replace('\xa0','').replace(' ','').replace(',','.')
    if re.fullmatch(r'-?\d+(\.\d+)?', t): return float(t)
    return None

wv=openpyxl.load_workbook(VOR, data_only=True, read_only=True)
vrows=[[c.value for c in r] for r in wv['Sheet1'].iter_rows()]
vor_by_num=defaultdict(list); vor_by_name=defaultdict(float); cur=None
for i,v in enumerate(vrows):
    head=s(v[0])+' '+s(v[1])
    m=re.search(r'ВОР\s*(\d\d-\d\d-\d\d)', head)
    if m and 'Ведомость' in head: cur=m.group(1); continue
    if cur and re.fullmatch(r'\d+', s(v[0])):
        rec=dict(row=i+1, name=s(v[1]), unit=s(v[2]), qty=N(v[3]), price=N(v[8]), total=N(v[9]))
        vor_by_num[(cur, s(v[0]))].append(rec)
        if rec['qty'] is not None: vor_by_name[(cur, key(v[1]))]+=rec['qty']

wi=openpyxl.load_workbook(ITOG, data_only=True, read_only=True)
irows=[[c.value for c in r] for r in wi['Спецификация'].iter_rows()]

def eq(a,b,tol=0.005):
    if a is None or b is None: return a is None and b is None
    return abs(a-b) <= max(tol, abs(b)*1e-6)

res=dict(line=[], aggr=[], qty_bad=[], price_bad=[], total_bad=[], nomatch=[], noqty=[])
for i,v in enumerate(irows):
    if i==0: continue
    a=s(v[0]); k=s(v[10]) if len(v)>10 else ''
    ma=re.fullmatch(r'ВОР\s*(\d+)', a); mk=re.search(r'ВОР\s*(\d\d-\d\d-\d\d)', k)
    if not (ma and mk): continue
    vn, nu = mk.group(1), ma.group(1)
    it=dict(row=i+1, name=s(v[1]), unit=s(v[5]), qty=N(v[6]), price=N(v[7]), total=N(v[8]), vor=vn, num=nu)
    cands=vor_by_num.get((vn,nu))
    if not cands: res['nomatch'].append(it); continue
    # выбираем кандидата с совпадающим наименованием, иначе первого
    c=next((x for x in cands if key(x['name'])==key(it['name'])), cands[0])
    q_agg=vor_by_name.get((vn, key(c['name'])))
    if   eq(it['qty'], c['qty']):  res['line'].append((it,c)); status='строка'
    elif eq(it['qty'], q_agg):     res['aggr'].append((it,c)); status='агрегат'
    else:
        res['qty_bad'].append((it,c,q_agg)); status='РАСХОЖДЕНИЕ'
    if it['price'] is not None and c['price'] is not None and not eq(it['price'], c['price'], 0.01):
        res['price_bad'].append((it,c))
    if c['price'] is None and it['price'] is not None: res['noqty'].append((it,c))

n=len(res['line'])+len(res['aggr'])+len(res['qty_bad'])
print('=== СВЕРКА ИТОГОВЫЙ ↔ ИСХОДНЫЙ ВОР (корректный разбор чисел) ===')
print('сопоставлено строк                       : %d' % n)
print('  кол-во = одной строке ВОРа             : %d' % len(res['line']))
print('  кол-во = сумме одноимённых строк ВОРа  : %d  (итоговый агрегирует — это нормально)' % len(res['aggr']))
print('  КОЛИЧЕСТВО НЕ СХОДИТСЯ                 : %d' % len(res['qty_bad']))
print('ссылка на ВОР есть, а строки в ВОРе нет  : %d' % len(res['nomatch']))
print('цена есть в итоговом, но нет в ВОРе      : %d' % len(res['noqty']))
print('ЦЕНА НЕ СХОДИТСЯ (там где есть в обоих)  : %d' % len(res['price_bad']))
print()
print('--- КОЛИЧЕСТВО НЕ СХОДИТСЯ ---')
print('%-7s %-11s %-5s %-40s %13s %13s %13s' % ('стр','ВОР','№','наименование','итоговый','строка ВОР','сумма по ВОР'))
for it,c,q in sorted(res['qty_bad'], key=lambda x: x[0]['row']):
    print('%-7d %-11s %-5s %-40s %13s %13s %13s' % (it['row'],it['vor'],it['num'],it['name'][:40],
          it['qty'], c['qty'], (round(q,4) if q is not None else '—')))
print()
print('--- ЦЕНА НЕ СХОДИТСЯ ---')
for it,c in sorted(res['price_bad'], key=lambda x:x[0]['row']):
    print('  стр.%-6d ВОР %-10s №%-5s %-40s итог=%-12s ВОР=%-12s' %
          (it['row'],it['vor'],it['num'],it['name'][:40],it['price'],c['price']))
print()
print('--- ССЫЛКА ЕСТЬ, СТРОКИ В ВОРе НЕТ ---')
for it in res['nomatch'][:20]:
    print('  стр.%-6d «ВОР %s, №%s» | %s' % (it['row'],it['vor'],it['num'],it['name'][:60]))
